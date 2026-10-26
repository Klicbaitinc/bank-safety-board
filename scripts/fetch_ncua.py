#!/usr/bin/env python3
"""Download NCUA insured-credit-union list and eight quarters of call reports.

Dollar amounts in the 5300 files are dollars, not thousands. The insured-list
workbook publishes return on average assets and the net worth ratio in percent.
"""
import argparse
import csv
import datetime as dt
import io
import json
import re
import zipfile
from collections import defaultdict
from pathlib import Path

import openpyxl

from formatutil import (
    assert_scale, nice_name, parse_date, parse_num, quarter_label, ratio,
    slugify, zip_code,
)
from http_util import USER_AGENT, get

ROOT = Path(__file__).resolve().parent.parent
DATA = ROOT / "data"
CACHE = DATA / "cache" / "ncua"
NCUA = "https://ncua.gov"
QUARTERLY_PAGE = NCUA + "/analysis/credit-union-corporate-call-report-data/quarterly-data"
LIST_PAGE = NCUA + "/analysis/credit-union-corporate-call-report-data"

MONTHS = {
    "january": 1, "february": 2, "march": 3, "april": 4, "may": 5, "june": 6,
    "july": 7, "august": 8, "september": 9, "october": 10, "november": 11, "december": 12,
}


def dump(path, obj):
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_suffix(path.suffix + ".tmp")
    tmp.write_text(json.dumps(obj, ensure_ascii=False, separators=(",", ":")))
    tmp.replace(path)


def download(url, dest, refresh):
    dest.parent.mkdir(parents=True, exist_ok=True)
    etag_path = dest.with_suffix(dest.suffix + ".etag")
    if dest.exists() and dest.stat().st_size > 0 and not refresh:
        print(f"cache {dest.name}", flush=True)
        return dest
    headers = {}
    if refresh and dest.exists() and etag_path.exists():
        headers["If-None-Match"] = etag_path.read_text().strip()
    print(f"GET {url}", flush=True)
    response = get(url, headers=headers, pause=1.0, timeout=300, stream=True, allow_status=(304,))
    if response.status_code == 304:
        print(f"  not modified {dest.name}", flush=True)
        response.close()
        return dest
    part = dest.with_suffix(dest.suffix + ".part")
    with part.open("wb") as handle:
        for chunk in response.iter_content(1024 * 256):
            if chunk:
                handle.write(chunk)
    response.close()
    part.replace(dest)
    etag = response.headers.get("ETag")
    if etag:
        etag_path.write_text(etag)
    return dest


def page_text(url, name, refresh):
    path = CACHE / name
    if path.exists() and not refresh:
        return path.read_text(errors="replace")
    response = get(url, pause=0.5, timeout=60)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(response.text)
    return response.text


def discover(refresh):
    quarterly = page_text(QUARTERLY_PAGE, "quarterly.html", refresh)
    found = sorted(set(re.findall(r"/files/publications/analysis/call-report-data-(\d{4})-(\d{2})\.zip", quarterly)))
    if len(found) < 4:
        raise SystemExit("could not find NCUA call-report zip links")
    quarters = [(int(year), int(month)) for year, month in found][-8:]
    urls = [
        f"{NCUA}/files/publications/analysis/call-report-data-{year:04d}-{month:02d}.zip"
        for year, month in quarters
    ]
    parent = page_text(LIST_PAGE, "call-report-index.html", refresh)
    lists = re.findall(r"/files/publications/analysis/federally-insured-credit-union-list-([a-z]+)-(\d{4})\.zip", parent)
    parsed = []
    for month_name, year in lists:
        month = MONTHS.get(month_name.lower())
        if month:
            parsed.append((int(year), month, month_name.lower()))
    if not parsed:
        raise SystemExit("could not find the federally insured credit union list")
    year, month, month_name = sorted(parsed)[-1]
    list_url = f"{NCUA}/files/publications/analysis/federally-insured-credit-union-list-{month_name}-{year}.zip"
    return urls, list_url


def read_zip_text(archive, name):
    try:
        raw = archive.read(name)
    except KeyError:
        # File names inside the zip vary in case.
        match = next(item for item in archive.namelist() if item.lower() == name.lower())
        raw = archive.read(match)
    for encoding in ("utf-8-sig", "latin-1"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("latin-1")


def header_index(header):
    return {cell.strip().lower(): i for i, cell in enumerate(header)}


def cell(row, idx, key):
    position = idx.get(key.lower())
    if position is None or position >= len(row):
        return None
    text = row[position].strip()
    return text or None


def dollars(value):
    number = parse_num(value)
    if number is None:
        return None
    return int(round(number))


def parse_fs220(text):
    reader = csv.reader(io.StringIO(text))
    header = next(reader)
    idx = header_index(header)
    out = {}
    for row in reader:
        charter = parse_num(cell(row, idx, "cu_number"))
        if charter is None:
            continue
        out[int(charter)] = {
            "assets": dollars(cell(row, idx, "acct_010")),
            "shares": dollars(cell(row, idx, "acct_018")),
            "loans": dollars(cell(row, idx, "acct_025b")),
            "delinquent": dollars(cell(row, idx, "acct_041b")),
            "members_call": dollars(cell(row, idx, "acct_083")),
        }
    return out


def parse_fs220a(text):
    reader = csv.reader(io.StringIO(text))
    header = next(reader)
    idx = header_index(header)
    out = {}
    for row in reader:
        charter = parse_num(cell(row, idx, "cu_number"))
        if charter is None:
            continue
        out[int(charter)] = {
            "net_worth": dollars(cell(row, idx, "acct_997")),
        }
    return out


def parse_branches(text):
    reader = csv.DictReader(io.StringIO(text))
    branches = defaultdict(int)
    corporate = defaultdict(int)
    hq = {}
    for row in reader:
        charter = parse_num(row.get("CU_NUMBER"))
        if charter is None:
            continue
        charter = int(charter)
        kind = (row.get("SiteTypeName") or "").strip()
        if kind == "Branch Office":
            branches[charter] += 1
        elif kind == "Corporate Office":
            corporate[charter] += 1
        if (row.get("MainOffice") or "").strip().lower() == "yes" and charter not in hq:
            hq[charter] = {
                "address": (row.get("PhysicalAddressLine1") or "").strip(),
                "city": nice_name((row.get("PhysicalAddressCity") or "").strip()),
                "state": (row.get("PhysicalAddressStateCode") or "").strip().upper(),
                "zip": zip_code(row.get("PhysicalAddressPostalCode")),
            }
    return branches, corporate, hq


def parse_trades(text):
    reader = csv.DictReader(io.StringIO(text))
    names = defaultdict(list)
    for row in reader:
        charter = parse_num(row.get("CU_NUMBER"))
        trade = (row.get("TradeName") or "").strip()
        if charter is None or not trade:
            continue
        bucket = names[int(charter)]
        if trade not in bucket:
            bucket.append(trade)
    return names


def parse_foicu(text):
    reader = csv.DictReader(io.StringIO(text))
    out = {}
    for row in reader:
        charter = parse_num(row.get("CU_NUMBER"))
        if charter is None:
            continue
        out[int(charter)] = {
            "insured_since": parse_date(row.get("INSURED_DATE")),
            "street": (row.get("STREET") or "").strip(),
            "cycle": parse_date(row.get("CYCLE_DATE")),
        }
    return out


def column(header, *needles):
    for index, name in enumerate(header):
        text = str(name or "").replace("\n", " ").lower()
        if all(needle in text for needle in needles):
            return index
    raise SystemExit(f"insured list is missing a column containing {needles}")


def load_insured_list(path):
    with zipfile.ZipFile(path) as archive:
        name = next(item for item in archive.namelist() if item.lower().endswith(".xlsx"))
        workbook = openpyxl.load_workbook(io.BytesIO(archive.read(name)), read_only=True, data_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    rows = sheet.iter_rows(values_only=True)
    header = list(next(rows))
    c_charter = column(header, "charter")
    c_quarter = column(header, "year", "quarter")
    c_name = column(header, "credit union name")
    c_street = column(header, "street")
    c_city = column(header, "city")
    c_state = column(header, "state")
    c_zip = column(header, "zip")
    c_type = column(header, "type")
    c_low = column(header, "low-income")
    c_members = column(header, "members")
    c_assets = column(header, "total assets")
    c_deposits = column(header, "total deposits")
    c_roa = column(header, "return on average assets")
    c_nwr = column(header, "net worth ratio")
    records = {}
    for row in rows:
        if not row or row[c_charter] is None:
            continue
        charter = parse_num(row[c_charter])
        if charter is None:
            continue
        quarter_value = parse_num(row[c_quarter])
        quarter = None
        if quarter_value is not None:
            year = int(quarter_value)
            qnum = int(round((quarter_value - year) * 10))
            month = {1: "03", 2: "06", 3: "09", 4: "12"}.get(qnum)
            if month:
                day = {"03": "31", "06": "30", "09": "30", "12": "31"}[month]
                quarter = f"{year}-{month}-{day}"
        official = str(row[c_name] or "").strip()
        low = str(row[c_low] or "").strip().lower()
        records[int(charter)] = {
            "charter": int(charter),
            "official_name": official,
            "name": nice_name(official),
            "slug": slugify(official),
            "address": str(row[c_street] or "").strip(),
            "city": nice_name(str(row[c_city] or "").strip()),
            "state": str(row[c_state] or "").strip().upper(),
            "zip": zip_code(row[c_zip]),
            "cu_type": str(row[c_type] or "").strip(),
            "low_income": True if low == "yes" else False if low == "no" else None,
            "members": dollars(row[c_members]),
            "list_assets": dollars(row[c_assets]),
            "list_deposits": dollars(row[c_deposits]),
            "roa": ratio(row[c_roa]),
            "net_worth_ratio": ratio(row[c_nwr]),
            "list_quarter": quarter,
        }
    workbook.close()
    return records


def cycle_from_name(url):
    match = re.search(r"call-report-data-(\d{4})-(\d{2})\.zip", url)
    year, month = int(match.group(1)), int(match.group(2))
    day = {3: 31, 6: 30, 9: 30, 12: 31}[month]
    return dt.date(year, month, day)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--refresh", action="store_true")
    args = parser.parse_args()
    print(f"NCUA user agent: {USER_AGENT}", flush=True)
    urls, list_url = discover(args.refresh)
    print("quarters:", ", ".join(urls), flush=True)
    print("insured list:", list_url, flush=True)

    list_path = download(list_url, CACHE / Path(list_url).name, args.refresh)
    insured = load_insured_list(list_path)
    print(f"insured credit unions {len(insured)}", flush=True)

    series = defaultdict(list)
    branches = {}
    corporate = {}
    hq = {}
    trades = {}
    foicu = {}
    latest_cycle = None
    for url in urls:
        cycle = cycle_from_name(url)
        path = download(url, CACHE / Path(url).name, args.refresh)
        with zipfile.ZipFile(path) as archive:
            fs220 = parse_fs220(read_zip_text(archive, "FS220.txt"))
            fs220a = parse_fs220a(read_zip_text(archive, "FS220A.txt"))
            if latest_cycle is None or cycle > latest_cycle:
                latest_cycle = cycle
                branches, corporate, hq = parse_branches(read_zip_text(archive, "Credit Union Branch Information.txt"))
                trades = parse_trades(read_zip_text(archive, "TradeNames.txt"))
                foicu = parse_foicu(read_zip_text(archive, "FOICU.txt"))
        print(f"  {cycle.isoformat()}: fs220 {len(fs220)} fs220a {len(fs220a)}", flush=True)
        for charter, row in fs220.items():
            extra = fs220a.get(charter) or {}
            series[charter].append({
                "date": cycle.isoformat(),
                "assets": row.get("assets"),
                "shares": row.get("shares"),
                "loans": row.get("loans"),
                "delinquent": row.get("delinquent"),
                "net_worth": extra.get("net_worth"),
            })

    credit_unions = []
    for charter, record in insured.items():
        record["quarters"] = sorted(series.get(charter, []), key=lambda item: item["date"])
        record["branches"] = branches.get(charter)
        record["corporate_offices"] = corporate.get(charter)
        record["trade_names"] = [nice_name(name) for name in trades.get(charter, [])]
        info = foicu.get(charter) or {}
        record["insured_since"] = info.get("insured_since")
        office = hq.get(charter) or {}
        if office.get("address"):
            record["hq_address"] = office.get("address") or ""
            record["hq_city"] = office.get("city") or record["city"]
            record["hq_state"] = office.get("state") or record["state"]
            record["hq_zip"] = office.get("zip") or record["zip"]
        else:
            record["hq_address"] = info.get("street") or record["address"]
            record["hq_city"] = record["city"]
            record["hq_state"] = record["state"]
            record["hq_zip"] = record["zip"]
        credit_unions.append(record)

    latest_rows = []
    for record in credit_unions:
        if record["quarters"] and record["quarters"][-1]["date"] == latest_cycle.isoformat():
            latest_rows.append(record)
    def computed_nwr(record):
        quarter = record["quarters"][-1]
        assets = quarter.get("assets")
        net_worth = quarter.get("net_worth")
        if not assets or net_worth is None:
            return None
        return net_worth / assets * 100
    def computed_dlq(record):
        quarter = record["quarters"][-1]
        loans = quarter.get("loans")
        delinquent = quarter.get("delinquent")
        if not loans or delinquent is None:
            return None
        return delinquent / loans * 100
    assert_scale("CU net worth ratio (call report)", [computed_nwr(r) for r in latest_rows], 4, 30)
    assert_scale("CU published net worth ratio", [r.get("net_worth_ratio") for r in latest_rows], 4, 30)
    assert_scale("CU published ROA", [r.get("roa") for r in latest_rows], -2, 5)
    assert_scale("CU delinquency ratio", [computed_dlq(r) for r in latest_rows], 0, 10)

    credit_unions.sort(key=lambda item: item["charter"])
    fetched = dt.datetime.now(dt.timezone.utc).replace(microsecond=0).isoformat().replace("+00:00", "Z")
    dump(DATA / "credit_unions.json", {
        "meta": {
            "source": "NCUA federally insured credit union list and 5300 call report",
            "fetched_at": fetched,
            "user_agent": USER_AGENT,
            "as_of": latest_cycle.isoformat(),
            "quarter": quarter_label(latest_cycle.isoformat()),
            "quarters": [cycle_from_name(url).isoformat() for url in urls],
            "insured_list": list_url,
            "call_reports": urls,
            "count": len(credit_unions),
        },
        "credit_unions": credit_unions,
    })
    print(f"wrote {len(credit_unions)} credit unions as of {latest_cycle.isoformat()}", flush=True)


if __name__ == "__main__":
    main()
